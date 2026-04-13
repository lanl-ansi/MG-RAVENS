from __future__ import annotations

import argparse
import ast
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_OUTDIR = ROOT / "out" / "smoke_autotemplate"
DEFAULT_WORKBOOK = DEFAULT_OUTDIR / "mentor_mismatch_review.xlsx"


SYNTHETIC_SUFFIXES = [
    "_anyOfPointer_anyOfContainer",
    "_anyOfContainer",
    "_anyOfPointer",
    "_PointerArray",
    "_Container",
    "_Array",
]


def _parse_listlike(value: Any) -> list[str]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    if isinstance(value, list):
        return [str(item) for item in value]
    text = str(value).strip()
    if not text or text == "[]":
        return []
    try:
        parsed = ast.literal_eval(text)
    except Exception:
        return [text]
    if isinstance(parsed, list):
        return [str(item) for item in parsed]
    return [text]


def _join_listlike(value: Any) -> str:
    items = _parse_listlike(value)
    return ", ".join(items)


def _base_name(name: str) -> str:
    base = str(name or "")
    changed = True
    while changed:
        changed = False
        for suffix in SYNTHETIC_SUFFIXES:
            if base.endswith(suffix):
                base = base[: -len(suffix)]
                changed = True
    return base


def _only_auto_kind(name: str) -> str:
    if name.endswith("_anyOfContainer"):
        return "Polymorphic wrapper"
    if name.endswith("_anyOfPointer"):
        return "Pointer union"
    if name.endswith("_PointerArray"):
        return "Array of pointers"
    if name.endswith("_Array"):
        return "Array wrapper"
    if name.endswith("_Container"):
        return "Container wrapper"
    return "Real class or special schema"


def _plain_missing_english(row: pd.Series) -> str:
    detail = str(row.get("classification_detail") or "")
    name = str(row.get("hand_name") or "")
    if detail == "real_class_absent_expected_inf_mkt_exclusion":
        return f"The hand template has {name}, but the auto build currently excludes it because it only lives on Inf/Mkt diagrams."
    if detail == "real_class_absent_from_auto_template":
        return f"The hand template has {name}, but the auto build did not create a schema for it."
    if detail == "synthetic_anyof_container":
        return f"The hand template has a polymorphic wrapper schema for {name}, but the auto build did not create that wrapper."
    if detail == "synthetic_anyof_pointer":
        return f"The hand template has a pointer-union schema for {name}, but the auto build did not create that union."
    if detail == "synthetic_pointer_array":
        return f"The hand template has an array-of-pointers schema for {name}, but the auto build did not create that array wrapper."
    if detail == "synthetic_array":
        return f"The hand template has an array wrapper schema for {name}, but the auto build did not create that wrapper."
    if detail == "synthetic_container":
        return f"The hand template has a container/hash-table wrapper schema for {name}, but the auto build did not create that wrapper."
    if detail == "synthetic_special":
        return f"The hand template has a special wrapper schema for {name}, but the auto build does not produce it."
    return f"The hand template has {name}, but the auto build does not currently match it."


def _current_read_missing(row: pd.Series) -> str:
    effective = str(row.get("effective_actionability") or "")
    diagram_status = str(row.get("diagram_exclusion_status") or "")
    if effective == "verified_hand_beyond_current_valid_diagrams":
        return "Best current read: the hand template is broader than the current valid UML diagrams here."
    if effective == "likely_hand_beyond_simplified_scope":
        return "Best current read: this probably belongs to hand-only or excluded-diagram scope."
    if effective == "likely_out_of_scope_for_current_diagrams":
        return "Best current read: this is probably outside the current simplified-diagram scope."
    if diagram_status == "not_on_any_simplified_diagram":
        return "Best current read: the current simplified diagrams do not show this at all."
    return "Best current read: this still needs a scope decision."


def _discussion_question_missing(row: pd.Series) -> str:
    detail = str(row.get("classification_detail") or "")
    effective = str(row.get("effective_actionability") or "")
    if effective == "verified_hand_beyond_current_valid_diagrams":
        return "Do we want to narrow the hand template, or expand the valid UML diagrams to support this?"
    if detail == "real_class_absent_expected_inf_mkt_exclusion":
        return "Should this stay excluded because it is on Inf/Mkt diagrams, or do we want those diagrams back in scope?"
    if detail == "real_class_absent_from_auto_template":
        return "Should this class stay out of scope, or should it be added to the valid UML/template pipeline?"
    return "Should we keep the hand shape, or align the hand template to the current UML-backed auto shape?"


def _plain_only_auto_english(name: str) -> str:
    kind = _only_auto_kind(name)
    base = _base_name(name)
    if kind == "Polymorphic wrapper":
        return f"The auto build created a polymorphic wrapper schema for {base}, but the hand template does not have the same wrapper."
    if kind == "Pointer union":
        return f"The auto build created a pointer-union schema for {base}, but the hand template does not have that union."
    if kind == "Array of pointers":
        return f"The auto build created an array-of-pointers schema for {base}, but the hand template does not have that wrapper."
    if kind == "Array wrapper":
        return f"The auto build created an array wrapper schema for {base}, but the hand template does not have that wrapper."
    if kind == "Container wrapper":
        return f"The auto build created a container/hash-table wrapper schema for {base}, but the hand template does not have that wrapper."
    return f"The auto build created a schema for {name}, but the hand template does not have a matching schema."


def _current_read_only_auto(name: str) -> str:
    kind = _only_auto_kind(name)
    if kind == "Real class or special schema":
        return "Best current read: the auto build is exposing a class or special shape that the hand template never modeled."
    return "Best current read: the auto build is keeping a wrapper shape that the hand template flattened or omitted."


def _discussion_question_only_auto(name: str) -> str:
    kind = _only_auto_kind(name)
    if kind == "Real class or special schema":
        return "Should the hand template gain this schema, or should auto be narrowed to hand scope?"
    return "Should the hand template gain this wrapper, or should auto collapse it to match hand behavior?"


def _problem_class_missing(row: pd.Series) -> str:
    name = str(row.get("hand_name") or "")
    return _base_name(name)


def _problem_class_only_auto(name: str) -> str:
    return _base_name(name)


def _problem_type_missing(row: pd.Series) -> str:
    detail = str(row.get("classification_detail") or "")
    return {
        "real_class_absent_from_auto_template": "Hand has a class that auto does not build",
        "real_class_absent_expected_inf_mkt_exclusion": "Hand has a class that auto excludes by policy",
        "synthetic_anyof_container": "Hand has a polymorphic wrapper that auto does not build",
        "synthetic_anyof_pointer": "Hand has a pointer union that auto does not build",
        "synthetic_pointer_array": "Hand has an array-of-pointers wrapper that auto does not build",
        "synthetic_array": "Hand has an array wrapper that auto does not build",
        "synthetic_container": "Hand has a container wrapper that auto does not build",
        "synthetic_special": "Hand has a special wrapper that auto does not build",
    }.get(detail, "Hand/auto mismatch")


def _problem_type_only_auto(name: str) -> str:
    kind = _only_auto_kind(name)
    return {
        "Polymorphic wrapper": "Auto builds a polymorphic wrapper that hand does not have",
        "Pointer union": "Auto builds a pointer union that hand does not have",
        "Array of pointers": "Auto builds an array-of-pointers wrapper that hand does not have",
        "Array wrapper": "Auto builds an array wrapper that hand does not have",
        "Container wrapper": "Auto builds a container wrapper that hand does not have",
        "Real class or special schema": "Auto builds a class or special schema that hand does not have",
    }[kind]


def _approach_to_fix_missing(row: pd.Series) -> str:
    detail = str(row.get("classification_detail") or "")
    effective = str(row.get("effective_actionability") or "")
    if detail == "real_class_absent_expected_inf_mkt_exclusion":
        return "Decide whether to keep Inf/Mkt diagrams out of scope. If yes, narrow the hand template. If no, bring those diagrams back into scope for auto."
    if effective in {
        "verified_hand_beyond_current_valid_diagrams",
        "likely_hand_beyond_simplified_scope",
        "likely_out_of_scope_for_current_diagrams",
    }:
        return "Treat this as a scope decision: either narrow the hand template to the current valid UML diagrams, or expand the UML/diagram scope if this behavior is still required."
    if detail == "real_class_absent_from_auto_template":
        return "If this class should really exist, add it to the valid UML/template pipeline. Otherwise remove or narrow it in the hand template."
    return "Either add the missing wrapper/shape to auto, or simplify the hand template if the wrapper is no longer wanted."


def _approach_to_fix_only_auto(name: str) -> str:
    kind = _only_auto_kind(name)
    if kind == "Real class or special schema":
        return "Decide whether the hand template should gain this schema, or whether auto should be narrowed so it no longer emits it."
    return "Decide whether to teach the hand template to keep this wrapper, or teach auto/schema composition to collapse it."


def _build_action_list(
    *,
    missing_review: pd.DataFrame,
    only_auto: pd.DataFrame,
) -> pd.DataFrame:
    missing_actions = pd.DataFrame(
        {
            "Problem Class": missing_review["Schema Name"].map(_base_name),
            "Schema Name": missing_review["Schema Name"],
            "Mismatch Direction": "Hand has it, auto is missing it",
            "Problem Type": missing_review.apply(_problem_type_missing, axis=1),
            "Suspected Problem": missing_review["Plain English"],
            "Best Current Read": missing_review["Current Read"],
            "Approach To Fix": missing_review.apply(_approach_to_fix_missing, axis=1),
            "Discussion Prompt": missing_review["Discuss With Mentor"],
            "Tracker Status": missing_review["Tracker Status"],
            "Diagrams": missing_review["Diagrams"],
            "Manual Note": missing_review["Manual Note"],
        }
    )

    only_auto_actions = pd.DataFrame(
        {
            "Problem Class": only_auto["Schema Name"].map(_problem_class_only_auto),
            "Schema Name": only_auto["Schema Name"],
            "Mismatch Direction": "Auto has it, hand is missing it",
            "Problem Type": only_auto["Schema Name"].map(_problem_type_only_auto),
            "Suspected Problem": only_auto["Plain English"],
            "Best Current Read": only_auto["Current Read"],
            "Approach To Fix": only_auto["Schema Name"].map(_approach_to_fix_only_auto),
            "Discussion Prompt": only_auto["Discuss With Mentor"],
            "Tracker Status": "Needs parity decision",
            "Diagrams": "",
            "Manual Note": "",
        }
    )

    action_list = pd.concat([missing_actions, only_auto_actions], ignore_index=True)
    action_list = action_list.sort_values(
        ["Tracker Status", "Problem Class", "Mismatch Direction", "Schema Name"],
        na_position="last",
    )
    return action_list


def _style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    note_fill = PatternFill("solid", fgColor="D9EAF7")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for column_cells in ws.columns:
            letter = get_column_letter(column_cells[0].column)
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 60)

        if ws.title == "Summary":
            for row in ws.iter_rows(min_row=2, max_col=2):
                if row[0].value == "Plain-English Readout":
                    row[0].fill = note_fill
                    row[1].fill = note_fill

    wb.save(path)


def build_workbook(outdir: Path = DEFAULT_OUTDIR, workbook_path: Path = DEFAULT_WORKBOOK) -> Path:
    outdir = Path(outdir)
    workbook_path = Path(workbook_path)
    temp_workbook_path = workbook_path.with_name(f"{workbook_path.stem}.__tmp__{workbook_path.suffix}")

    summary = json.loads((outdir / "summary.json").read_text())
    missing = pd.read_csv(outdir / "missing_classification.csv")
    normalized = pd.read_csv(outdir / "schema_keys_normalized.csv")
    gaps = pd.read_csv(outdir / "template_shape_gaps.csv")

    gap_cols = [
        "hand_name",
        "shape_mismatch_kind",
        "scope_status",
        "effective_actionability",
        "manual_review_status",
        "manual_review_note",
        "manual_future_followup",
        "hand_example_paths",
        "auto_example_paths",
    ]
    missing_review = missing.merge(gaps[gap_cols], how="left", on="hand_name")
    missing_review["Diagrams"] = missing_review["diagram_names"].map(_join_listlike)
    missing_review["Hand Example Paths"] = missing_review["hand_example_paths"].map(_join_listlike)
    missing_review["Auto Example Paths"] = missing_review["auto_example_paths"].map(_join_listlike)
    missing_review["Plain English"] = missing_review.apply(_plain_missing_english, axis=1)
    missing_review["Current Read"] = missing_review.apply(_current_read_missing, axis=1)
    missing_review["Discuss With Mentor"] = missing_review.apply(_discussion_question_missing, axis=1)
    missing_review["Manual Note"] = missing_review["manual_review_note"].fillna("")
    missing_review["Future Follow-up"] = missing_review["manual_future_followup"].fillna("")
    missing_review["Kind"] = missing_review["classification_detail"].map(
        {
            "real_class_absent_from_auto_template": "Real class missing from auto",
            "real_class_absent_expected_inf_mkt_exclusion": "Expected Inf/Mkt exclusion",
            "synthetic_anyof_container": "Missing polymorphic wrapper",
            "synthetic_anyof_pointer": "Missing pointer union",
            "synthetic_pointer_array": "Missing pointer array",
            "synthetic_array": "Missing array wrapper",
            "synthetic_container": "Missing container wrapper",
            "synthetic_special": "Missing special wrapper",
        }
    ).fillna(missing_review["classification_detail"])

    missing_review = missing_review[
        [
            "hand_name",
            "Kind",
            "Plain English",
            "Current Read",
            "Discuss With Mentor",
            "effective_actionability",
            "diagram_exclusion_status",
            "Diagrams",
            "Manual Note",
            "Future Follow-up",
            "Hand Example Paths",
            "Auto Example Paths",
        ]
    ].rename(
        columns={
            "hand_name": "Schema Name",
            "effective_actionability": "Tracker Status",
            "diagram_exclusion_status": "Diagram Scope Flag",
        }
    )
    missing_review = missing_review.sort_values(["Tracker Status", "Kind", "Schema Name"], na_position="last")

    only_auto = normalized[normalized["status"] == "only_in_auto"].copy()
    only_auto["Schema Name"] = only_auto["matched_auto_name"]
    only_auto["Base Name"] = only_auto["Schema Name"].map(_base_name)
    only_auto["Kind"] = only_auto["Schema Name"].map(_only_auto_kind)
    only_auto["Plain English"] = only_auto["Schema Name"].map(_plain_only_auto_english)
    only_auto["Current Read"] = only_auto["Schema Name"].map(_current_read_only_auto)
    only_auto["Discuss With Mentor"] = only_auto["Schema Name"].map(_discussion_question_only_auto)
    only_auto = only_auto[
        ["Schema Name", "Base Name", "Kind", "Plain English", "Current Read", "Discuss With Mentor"]
    ].sort_values(["Kind", "Schema Name"])

    manual = gaps[gaps["manual_review_status"].notna()].copy()
    manual["Diagrams"] = manual["diagram_names"].map(_join_listlike)
    manual = manual[
        [
            "hand_name",
            "effective_actionability",
            "manual_review_status",
            "manual_review_source",
            "manual_review_note",
            "manual_future_followup",
            "Diagrams",
        ]
    ].rename(
        columns={
            "hand_name": "Schema Name",
            "effective_actionability": "Tracker Status",
            "manual_review_status": "Manual Review Status",
            "manual_review_source": "Manual Review Source",
            "manual_review_note": "Plain-English Note",
            "manual_future_followup": "Suggested Follow-up",
        }
    ).sort_values(["Tracker Status", "Schema Name"])

    only_auto_counts = only_auto["Kind"].value_counts().rename_axis("Kind").reset_index(name="Count")
    missing_counts = missing_review["Kind"].value_counts().rename_axis("Kind").reset_index(name="Count")
    action_list = _build_action_list(missing_review=missing_review, only_auto=only_auto)

    summary_rows = [
        ("Hand schemas", summary["hand_schema_count"]),
        ("Auto schemas", summary["auto_schema_count"]),
        ("Shared exact matches", summary["normalized_schema_key_counts"]["shared_exact"]),
        ("Missing in auto", summary["normalized_schema_key_counts"]["missing_in_auto"]),
        ("Only in auto", summary["normalized_schema_key_counts"]["only_in_auto"]),
        ("Targeted regression failures", f"{summary['targeted_check_failures']} / {summary['targeted_check_total']}"),
        (
            "Plain-English Readout",
            "The tracker currently shows no open builder gaps. The remaining differences are scope and parity questions between the hand template and the current valid simplified UML diagrams.",
        ),
    ]
    summary_df = pd.DataFrame(summary_rows, columns=["Metric", "Value"])

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(temp_workbook_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        action_list.to_excel(writer, sheet_name="Action_List", index=False)
        missing_counts.to_excel(writer, sheet_name="Missing_Counts", index=False)
        only_auto_counts.to_excel(writer, sheet_name="Only_Auto_Counts", index=False)
        missing_review.to_excel(writer, sheet_name="Missing_in_Auto", index=False)
        only_auto.to_excel(writer, sheet_name="Only_in_Auto", index=False)
        manual.to_excel(writer, sheet_name="Manual_Rulings", index=False)

    _style_workbook(temp_workbook_path)
    try:
        temp_workbook_path.replace(workbook_path)
        return workbook_path
    except PermissionError:
        fallback_workbook_path = workbook_path.with_name(f"{workbook_path.stem}_updated{workbook_path.suffix}")
        if fallback_workbook_path.exists():
            fallback_workbook_path.unlink()
        temp_workbook_path.replace(fallback_workbook_path)
        return fallback_workbook_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Export a mentor-friendly mismatch review workbook.")
    parser.add_argument("--outdir", type=Path, default=DEFAULT_OUTDIR)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()

    workbook_path = build_workbook(outdir=args.outdir, workbook_path=args.workbook)
    print(workbook_path)


if __name__ == "__main__":
    main()
